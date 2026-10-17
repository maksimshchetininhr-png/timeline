"""Import data, never instructions, from the two source workbooks.
Usage: python3 scripts/import_workbooks.py /path/to/master.xlsx /path/to/matrix.xlsx
"""
import json, sys, re, os
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parent.parent
paths=list(Path('/workspace/attachments').rglob('*.xlsx'))
master=Path(sys.argv[1]) if len(sys.argv)>1 else next(p for p in paths if p.name=='life_timeline_master.xlsx')
matrix=Path(sys.argv[2]) if len(sys.argv)>2 else next(p for p in paths if p.name=='life_timeline_matrix_master.xlsx')
def date(v): return v.date().isoformat() if isinstance(v,datetime) else None
countries={'Russia':(61,95),'Netherlands':(52.2,5.3),'USA':(39,-99),'Norway':(62,10),'Finland':(64,26),'Sweden':(62,15),'Spain':(40,-3),'Estonia':(59,25),'Czechia':(49.8,15.5),'Germany':(51,10),'Poland':(52,19),'Austria':(47.5,14.5),'Slovakia':(48.7,19.7),'Italy':(42.5,12.5),'UK':(54,-2),'Maldives':(3.2,73.2),'Hungary':(47,19),'Belgium':(50.7,4.5),'Portugal':(39.5,-8),'China':(35,104),'Serbia':(44,21),'France':(47,2),'Kazakhstan':(48,68)}
cities={'Murmansk':(68.97,33.07,'Russia'),'Severomorsk':(69.07,33.42,'Russia'),'Moscow':(55.76,37.62,'Russia'),'Korolev':(55.92,37.85,'Russia'),'Amsterdam':(52.37,4.9,'Netherlands'),'Haarlem':(52.39,4.64,'Netherlands'),'Sciacca':(37.51,13.08,'Italy'),'San Francisco':(37.77,-122.42,'USA'),'Stockholm':(59.33,18.07,'Sweden'),'Prague':(50.08,14.43,'Czechia'),'Berlin':(52.52,13.41,'Germany'),'Vienna':(48.21,16.37,'Austria'),'Bratislava':(48.15,17.11,'Slovakia'),'Tallinn':(59.44,24.75,'Estonia'),'Helsinki':(60.17,24.94,'Finland'),'Dresden':(51.05,13.74,'Germany'),'Krakow':(50.06,19.94,'Poland'),'Oświęcim':(50.04,19.22,'Poland'),'Kirkenes':(69.73,30.05,'Norway'),'London':(51.51,-0.13,'UK'),'Saint Petersburg':(59.93,30.34,'Russia'),'Tuapse':(44.1,39.07,'Russia'),'Balkhash':(46.85,74.98,'Kazakhstan'),'Alma-Ata':(43.24,76.95,'Kazakhstan'),'Nyrob':(60.73,56.72,'Russia')}
def place(raw,home=False):
 if not raw: return None
 # A route is a set of possible cities, not an exact location for every trip day.
 country=next((c for c in countries if c in raw),None)
 multi=';' in raw or ' and ' in raw
 city=next((c for c in cities if c in raw),None)
 inferred=False
 if home and not city and any(x in raw for x in ['Zoi','Beringa','Baumana','Karla Libknekhta']): city='Murmansk';inferred=True
 if city and not country: country=cities[city][2]
 if multi:
  city=None
 coords=cities[city][:2] if city else countries.get(country)
 return {'label': (city+', '+country) if city else (country or raw), 'city':city,'country':country,'coordinates':list(reversed(coords)) if coords else None,'route':raw if multi else None,'inferred':inferred,'precision':'route' if multi else 'city' if city else 'country'}
w=load_workbook(master,data_only=True)
events=[]
for row_num,row in enumerate(list(w['Events'].values)[1:],2):
 start,end,year,month,category,sub,title,loc,person,certainty,notes=row
 if not start: continue
 owner='both' if person=='Nadya' else 'family' if category=='Family' else 'maksim'
 cat={'Personal':'Life','Relocation':'Home','Housing':'Home','Property':'Home','Side hustle':'Career','Side work':'Career','Car':'Life','Certification':'Milestones','COVID / restrictions':'World'}.get(category,category)
 if sub=='Relationship':cat='Together'
 if sub=='Wedding' and person=='Nadya':cat='Together'
 p=place(loc,home=category=='Housing')
 if category=='Career' and person: title=f'{title} · {person}'
 if category=='Family' and person=='Dad' and not title.lower().startswith('dad'): title='Dad · '+title
 if title=='Residence / move':title='A new home'+(' in '+p['label'].split(',')[0] if p else '')
 if title=='Born':title='Maksim is born'
 if title=='Wedding with Nadya':title='We got married'
 if title=='First date with Nadya':title='Our first date'
 if title=='Moved in together with Nadya':title='A place of our own, together'
 # Financial sheets are outside the family-story scope. Do not expose financial notes.
 clean_notes='; '.join(s.strip() for s in (notes or '').split(';') if not re.search(r'€|salary|income|\bRUB\b|net/month|RSU',s,re.I))
 events.append({'id':f'import-{row_num}','start':date(start),'end':date(end),'title':title,'category':cat,'originalCategory':category,'subcategory':sub,'person':owner,'place':p,'certainty':certainty or 'Unknown','notes':clean_notes,'source':f'{master.name} · Events · row {row_num}','featured':owner=='both' or category=='Relocation' or (category=='Property' and 'Haarlem' in (loc or '')),'residence':category=='Housing' and sub=='Residence','trip':category=='Travel' and sub not in ['Residence/country timeline']})
# Include the matrix as dated context, without carrying a blank cell forward.
mw=load_workbook(matrix,data_only=True)
rows=list(mw.active.values)
context=[]
for row in rows[1:]:
 if row[0] not in ['Education','Career','Home / location','Family']:continue
 for i,(dt,value) in enumerate(zip(rows[0][1:],row[1:])):
  if value is None or not isinstance(dt,datetime):continue
  following=rows[0][i+2] if i+2<len(rows[0]) else None
  context.append({'category':row[0],'start':date(dt),'end':date(following),'value':str(value)})
result={'events':events,'context':context,'sources':[master.name,matrix.name],'importedAt':datetime.now().date().isoformat(),'people':[{'id':'maksim','name':'Maksim'},{'id':'nadya','name':'Nadya'}]}
output=Path(os.environ.get('TIMELINE_DATA_FILE', str(Path(os.environ.get('DATA_DIR', str(ROOT/'.local')))/'timeline.json')))
output.parent.mkdir(parents=True,exist_ok=True)
output.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
output.chmod(0o600)
print(f'Imported {len(events)} events and {len(context)} context entries from both workbooks.')
